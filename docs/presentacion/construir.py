"""Genera la presentacion del taller: "Pruebas de UI y UX.pptx" en la raiz del repositorio.

    python docs/presentacion/construir.py            (desde la raiz del repo)

Usa el mismo sistema visual que la presentacion del taller de carga (kit.py) y
las capturas reales de docs/presentacion/img, que produce capturas.js. Si la
interfaz cambia, primero se repiten las capturas y despues se regenera esto.

Todas las cifras de las diapositivas salen del repositorio (README, guias,
comentarios de las pruebas y reportes del modulo 6). Si una cambia alla, hay
que cambiarla aqui: la presentacion no las lee sola.
"""
import os
import sys

from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

sys.path.insert(0, os.path.dirname(__file__))
from kit import (C, Deck, W, H, FONT, MONO, text, para_space, box, shape_text, line, code,
                 Plot, cat, pill, dot, rgb)

AQUI = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.abspath(os.path.join(AQUI, '..', '..'))
IMG = os.path.join(AQUI, 'img')
SNAP = os.path.join(RAIZ, 'playwright', 'tests', 'modulo4-visual.spec.js-snapshots')
SALIDA = os.path.join(RAIZ, 'Pruebas de UI y UX.pptx')

d = Deck(pie='Taller de pruebas de UI y UX')


# ---------------------------------------------------------------- piezas propias

def navegador(s, x, y, w, img, url='localhost:8080/', max_h=None):
    """Captura real dentro de un marco de navegador. Devuelve el borde inferior."""
    from PIL import Image
    iw, ih = Image.open(img).size
    barra = 0.34
    h_img = w * ih / iw
    recorte = 0
    if max_h and h_img > max_h - barra:
        recorte = 1 - (max_h - barra) / h_img
        h_img = max_h - barra
    marco = box(s, x, y, w, barra + h_img, fill='#ffffff', line=C['border'], radius=0.02, line_w=1)
    box(s, x + 0.01, y + 0.01, w - 0.02, barra - 0.01, fill=C['plane'], radius=0.0,
        shape=MSO_SHAPE.RECTANGLE)
    for i, col in enumerate([C['critical'], C['warning'], C['good']]):
        dot(s, x + 0.2 + i * 0.17, y + barra / 2, 0.1, col)
    u = box(s, x + 0.78, y + 0.07, min(w - 0.95, 3.2), barra - 0.14, fill='#ffffff', radius=0.5)
    shape_text(u, url, size=9, color=C['ink2'], align=PP_ALIGN.LEFT)
    pic = s.shapes.add_picture(img, Inches(x), Inches(y + barra), width=Inches(w))
    if recorte:
        pic.crop_bottom = recorte
        pic.height = Inches(h_img)
    return y + barra + h_img


def tarjeta(s, x, y, w, h, fill=None, line_=None):
    return box(s, x, y, w, h, fill=fill or C['plane'], line=line_, radius=0.05)


def flecha(s, x1, y1, x2, y2, color=None):
    return line(s, x1, y1, x2, y2, color=color or C['axis'], width=1.5, arrow_end=True)


def estado(s, x, y, etiqueta, tipo, w=1.55):
    """Pastilla de veredicto: detectado / sobrevivio / no ejercitado."""
    col, fondo = {'si': (C['goodtext'], '#d7efd7'), 'no': (C['critical'], '#f6d8d8'),
                  'aviso': ('#7a5600', '#fdecc4')}[tipo]
    return pill(s, x, y, etiqueta, col, fondo, size=11, w=w)


# ---------------------------------------------------------------- 1. portada

s = d.slide(notes=(
    'Presentación del taller. La idea que ordena toda la sesión está en el subtítulo: '
    'hay preguntas que se responden ejecutando código (¿la interfaz funciona?) y una que solo '
    'se responde observando personas (¿la gente entiende cómo usarla?). '
    'La captura es la aplicación real que van a probar, la misma Registraduría de los '
    'talleres de unitarias, integración y carga, ahora con interfaz web.'))
text(s, 0.6, 1.2, 6.8, 0.3, 'TESTING Y VALIDACIÓN DE SOFTWARE', size=12, color=C['s1'], bold=True, spacing=1.5)
text(s, 0.6, 1.6, 6.8, 1.9, 'Pruebas de UI y UX', size=48, bold=True, line=1.0)
text(s, 0.6, 2.55, 6.6, 1.3,
     '¿La interfaz funciona? Eso lo responde una máquina. ¿La gente entiende cómo usarla? '
     'Eso solo se responde mirando a personas.', size=20, color=C['ink2'], line=1.15)
px, fy = 0.6, 4.05
for et in ['E2E', 'Page Object Model', 'Accesibilidad', 'Regresión visual', 'Usabilidad', 'Probar con IA']:
    ancho = 0.105 * len(et) + 0.45
    if px + ancho > 7.2:
        px, fy = 0.6, fy + 0.45
    pill(s, px, fy, et, C['blue600'], C['blue100'], size=11, w=ancho)
    px += ancho + 0.12
text(s, 0.6, 5.55, 6.6, 0.9, [
    [('Sistema bajo prueba: ', {'bold': True, 'color': C['ink']}),
     ('la Registraduría (Spring Boot), la misma aplicación de los talleres de pruebas unitarias, '
      'de integración y de carga.', {})]], size=13, color=C['ink2'], line=1.2)
navegador(s, 7.75, 0.8, 5.0, os.path.join(IMG, 'portada.png'), max_h=5.95)


# ---------------------------------------------------------------- 2. mapa

s = d.slide('Siete módulos, cinco preguntas distintas', 'Mapa de la sesión', notes=(
    'Mapa del taller. Los módulos 1 a 4 se automatizan con Playwright (28 pruebas en total; '
    'Selenium es la pista alternativa, con 7). El 5 no se automatiza: es una sesión con cinco '
    'personas. El 6 es exploratorio y no suma puntos. '
    'Conviene decir desde el principio que el 3B existe para ver fallar a la herramienta: '
    'el módulo 3 sale verde y eso, como aprendizaje, no enseña a leer un informe de axe.'))
mods = [
    ('1', 'Pruebas E2E', '¿La interfaz funciona?', '9 pruebas', 'auto'),
    ('2', 'Page Object Model', '¿Las pruebas sobreviven a un rediseño?', '5 pruebas', 'auto'),
    ('3', 'Accesibilidad', '¿Cumple lo que una máquina puede medir de WCAG?', '5 pruebas', 'auto'),
    ('3B', 'Ver fallar a axe', '¿Qué no ve axe? 17 defectos sembrados', '6 pruebas', 'auto'),
    ('4', 'Regresión visual', '¿Cambió algo sin querer?', '3 pruebas', 'auto'),
    ('5', 'Usabilidad con usuarios', '¿La gente entiende cómo usarla?', 'sin automatizar', 'persona'),
    ('6', 'Probar con IA', '¿Sirve una IA para probar? Se mide, no se opina', 'exploratorio · 0 pts', 'ia'),
    ('→', 'Entrega', 'Repositorio, Wiki, CI y reflexión final', 'rúbrica de 65 pts', 'entrega'),
]
cw, ch_, gx = 2.83, 2.35, 0.27
for i, (num, nom, preg, cuenta, tipo) in enumerate(mods):
    cx = 0.6 + (i % 4) * (cw + gx)
    cy = 1.75 + (i // 4) * (ch_ + 0.3)
    fondo = {'auto': C['plane'], 'persona': C['orange100'], 'ia': C['plane'], 'entrega': '#ffffff'}[tipo]
    tarjeta(s, cx, cy, cw, ch_, fill=fondo, line_=C['border'] if tipo == 'entrega' else None)
    text(s, cx + 0.25, cy + 0.2, 1.2, 0.6, num, size=28, bold=True,
         color=C['s2'] if tipo == 'persona' else C['s1'])
    text(s, cx + 0.25, cy + 0.85, cw - 0.45, 0.4, nom, size=15, bold=True)
    text(s, cx + 0.25, cy + 1.2, cw - 0.45, 0.7, preg, size=12, color=C['ink2'], line=1.15)
    col, fnd = {'auto': (C['blue600'], C['blue100']), 'persona': (C['s2'], '#ffffff'),
                'ia': (C['s7'], '#e4e0f6'), 'entrega': (C['ink2'], C['plane'])}[tipo]
    pill(s, cx + 0.25, cy + ch_ - 0.5, cuenta, col, fnd, size=10, w=0.085 * len(cuenta) + 0.4)


# ---------------------------------------------------------------- 3. UI y UX

s = d.slide('Automatizar responde cuatro preguntas. La quinta, no', 'UI y UX', notes=(
    'La tesis del taller. Las cuatro primeras técnicas se automatizan y responden si la '
    'interfaz funciona, si las pruebas aguantan cambios, si cumple lo medible de WCAG y si '
    'algo cambió sin querer. Ninguna dice si una persona entiende la pantalla. '
    'Accesibilidad queda en la frontera a propósito: axe mide cerca del 40 %, el resto es juicio humano.'))
tecs = [
    ('E2E', '¿La interfaz funciona?', 'M1'),
    ('Page Objects', '¿Las pruebas sobreviven a un rediseño?', 'M2'),
    ('Accesibilidad', '¿Cumple lo que una máquina puede medir de WCAG?', 'M3 · 3B'),
    ('Regresión visual', '¿Cambió algo sin querer?', 'M4'),
    ('Usabilidad', '¿La gente entiende cómo usarla?', 'M5'),
]
tw, tg, ty = 2.25, 0.22, 2.5
for i, (nom, preg, m) in enumerate(tecs):
    tx = 0.6 + i * (tw + tg)
    persona = i == 4
    tarjeta(s, tx, ty, tw, 2.2, fill=C['orange100'] if persona else C['plane'])
    text(s, tx + 0.22, ty + 0.22, tw - 0.4, 0.3, m, size=11, bold=True,
         color=C['s2'] if persona else C['s1'], spacing=1)
    text(s, tx + 0.22, ty + 0.55, tw - 0.4, 0.4, nom, size=16, bold=True)
    text(s, tx + 0.22, ty + 1.0, tw - 0.4, 1.0, preg, size=13, color=C['ink2'], line=1.15)
# llaves: que responde una maquina y que solo una persona
x_fin_auto = 0.6 + 4 * tw + 3 * tg
line(s, 0.6, 2.2, x_fin_auto, 2.2, color=C['s1'], width=1.5)
line(s, 0.6, 2.2, 0.6, 2.35, color=C['s1'], width=1.5)
line(s, x_fin_auto, 2.2, x_fin_auto, 2.35, color=C['s1'], width=1.5)
text(s, 0.6, 1.72, 8, 0.4, 'UI · lo mide una máquina ejecutando pruebas', size=13, bold=True, color=C['s1'])
x5 = 0.6 + 4 * (tw + tg)
line(s, x5, 2.2, x5 + tw, 2.2, color=C['s2'], width=1.5)
line(s, x5, 2.2, x5, 2.35, color=C['s2'], width=1.5)
line(s, x5 + tw, 2.2, x5 + tw, 2.35, color=C['s2'], width=1.5)
text(s, x5, 1.72, tw + 0.3, 0.4, 'UX · observando personas', size=13, bold=True, color=C['s2'])
text(s, 0.6, 5.25, 11.5, 1.2, [
    [('Un formulario puede pasar las 28 pruebas automatizadas ', {'bold': True, 'color': C['ink']}),
     ('y aun así hacer que la gente abandone en el segundo campo. La usabilidad se mide observando '
      'a personas, no ejecutando código.', {})]], size=18, color=C['ink2'], line=1.2)


# ---------------------------------------------------------------- 4. sistema bajo prueba

s = d.slide('La regla de negocio vive en dos lugares', 'Sistema bajo prueba', notes=(
    'Recorrer el camino de una inscripción. El navegador valida primero (nombre, documento, '
    'edad entre 0 y 120) y solo si pasa llama a POST /register. El servidor vuelve a validar '
    'y además conoce lo que el navegador no puede saber: si la persona está viva y si el '
    'documento ya existe. Esta duplicación es la que explota la prueba 09 más adelante. '
    'La captura es un rechazo real por minoría de edad.'))
pasos = [
    ('Persona en el navegador', 'Llena el formulario y pulsa "Registrar votante".', C['plane'], C['ink']),
    ('Capa 1 · app.js en el navegador', 'Valida nombre, documento y edad entre 0 y 120. '
     'Si falla, muestra el error y no llama al servidor.', C['blue100'], C['blue600']),
    ('Capa 2 · POST /register (Spring Boot)', 'Valida documento, vida, edad y duplicados, en ese orden, '
     'y responde VALID o el motivo: INVALID, DEAD, INVALID_AGE, UNDERAGE o DUPLICATED.', C['blue100'], C['blue600']),
    ('Resultado en pantalla', 'Un mensaje con role="status", para que un lector de pantalla lo anuncie.',
     C['plane'], C['ink']),
]
py_ = 1.75
for i, (tit, desc, fondo, col) in enumerate(pasos):
    tarjeta(s, 0.6, py_, 6.9, 1.0, fill=fondo)
    text(s, 0.85, py_ + 0.14, 6.4, 0.35, tit, size=15, bold=True, color=col)
    text(s, 0.85, py_ + 0.48, 6.4, 0.5, desc, size=12, color=C['ink2'], line=1.1)
    if i < len(pasos) - 1:
        flecha(s, 1.2, py_ + 1.0, 1.2, py_ + 1.27)
    py_ += 1.27
navegador(s, 8.5, 1.55, 3.45, os.path.join(IMG, 'formulario-resultado.png'))


# ---------------------------------------------------------------- 5. localizadores

s = d.slide('Buscar los elementos como los busca una persona', 'Módulo 1 · Localizadores', notes=(
    'Tres formas de encontrar el mismo botón. getByRole sobrevive a un rediseño y además '
    'comprueba que el botón tiene nombre accesible: si un lector de pantalla no lo encuentra, '
    'la prueba tampoco. Las clases CSS cambian con el primer ajuste de estilos, y un XPath '
    'absoluto se rompe si alguien agrega un div. La rúbrica penaliza el XPath absoluto.'))
filas = [
    ('Bien', C['goodtext'], '#d7efd7', "page.getByRole('button', { name: 'Registrar votante' })",
     'Sobrevive a un rediseño y, de paso, comprueba que el botón es accesible.'),
    ('Mal', '#7a5600', '#fdecc4', "page.locator('.btn.btn-primary.mt-3')",
     'Se rompe con el primer cambio de estilos.'),
    ('Peor', C['critical'], '#f6d8d8', "page.locator('/html/body/div[2]/form/button')",
     'XPath absoluto: se rompe si alguien agrega un div. Lo penaliza la rúbrica.'),
]
for i, (et, col, fnd, src, por) in enumerate(filas):
    fy = 1.85 + i * 1.5
    pill(s, 0.6, fy + 0.2, et, col, fnd, size=12, w=0.95)
    code(s, 1.8, fy, 6.7, 0.72, src, size=14)
    text(s, 8.85, fy - 0.05, 3.9, 0.82, por, size=14, color=C['ink2'], line=1.2, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.6, 6.4, 12, 0.4, 'Orden de preferencia de la guía: rol → etiqueta → texto → test id → CSS o XPath, como último recurso.',
     size=12, color=C['muted'])


# ---------------------------------------------------------------- 6. esperas y la prueba inestable

s = d.slide('Una prueba inestable casi nunca es mala suerte', 'Módulo 1 · Esperas', notes=(
    'Dos lecciones sobre estabilidad. La primera: cero esperas fijas; Playwright reintenta la '
    'aserción hasta que se cumple o se acaba el tiempo. La segunda es una historia real de este '
    'repositorio: el documento único se generaba con Date.now() % 1000000. Parece único y no lo es. '
    'Con --repeat-each=30 y 8 workers, dos copias arrancaban en el mismo milisegundo y una de cada '
    '30 fallaba con "Documento ya inscrito". No era el servidor ni la red: era la prueba.'))
text(s, 0.6, 1.7, 5.9, 0.4, 'Cero esperas fijas', size=18, bold=True)
code(s, 0.6, 2.2, 6.1, 1.85,
     "// Una apuesta: lento si responde rápido,\n// insuficiente si responde lento\n"
     "await page.waitForTimeout(3000);\n\n// Reintenta hasta que se cumple o expira\n"
     "await expect(page.getByRole('heading',\n  { name: 'Inscripción exitosa' })).toBeVisible();",
     size=12, highlight={2: C['critical'], 5: C['goodtext'], 6: C['goodtext']})
text(s, 0.6, 4.35, 6.1, 1.8,
     'Un sleep es la causa número uno de pruebas inestables. La entrega exige cero, y la rúbrica '
     'las penaliza junto con el XPath absoluto.', size=14, color=C['ink2'], line=1.2)
tarjeta(s, 7.2, 1.7, 5.55, 4.85, fill='#ffffff', line_=C['border'])
text(s, 7.5, 1.95, 5.0, 0.3, 'MEDIDO EN ESTE REPOSITORIO', size=11, bold=True, color=C['s2'], spacing=1)
text(s, 7.5, 2.3, 5.0, 1.0, '1 de cada 30', size=44, bold=True, color=C['critical'])
text(s, 7.5, 3.2, 5.0, 0.7, ['ejecuciones fallaba con "Documento ya inscrito".', 'Medido con --repeat-each=30 y 8 workers.'],
     size=13, color=C['ink2'], line=1.15)
code(s, 7.5, 4.05, 4.95, 1.05,
     "// Antes: dos copias, mismo milisegundo\nDate.now() % 1000000\n"
     "// Ahora: 900 millones de valores\nMath.floor(Math.random() * 900_000_000)",
     size=11, highlight={1: C['critical'], 3: C['goodtext']})
text(s, 7.5, 5.35, 5.0, 1.0, 'Para cazar una prueba inestable, repítala en paralelo: --repeat-each es la forma estándar.',
     size=12, color=C['ink2'], line=1.15)


# ---------------------------------------------------------------- 7. POM

s = d.slide('Las pruebas se leen como reglas de negocio', 'Módulo 2 · Page Object Model', notes=(
    'El mismo escenario del módulo 1 escrito con un Page Object. La prueba dice qué se está '
    'probando (una persona de 17 años es rechazada) y RegistroPage sabe cómo se interactúa con la '
    'pantalla. Si el botón cambia de texto, se corrige en un lugar. En la entrega, ninguna prueba '
    'puede tener un localizador directo.'))
text(s, 0.6, 1.7, 5.8, 0.4, 'Sin Page Object: una secuencia de clics', size=15, bold=True, color=C['ink2'])
code(s, 0.6, 2.15, 5.9, 2.2,
     "await page.getByLabel('Nombre completo').fill('Ana');\n"
     "await page.getByLabel('Número de documento')\n  .fill(String(documento));\n"
     "await page.getByLabel('Edad').fill('17');\n"
     "await page.getByRole('button',\n  { name: 'Registrar votante' }).click();\n"
     "await expect(page.getByRole('heading',\n  { name: 'Persona menor de edad' })).toBeVisible();",
     size=11)
text(s, 6.9, 1.7, 5.8, 0.4, 'Con Page Object: la regla que se prueba', size=15, bold=True, color=C['s1'])
code(s, 6.9, 2.15, 5.85, 1.0,
     "await registro.inscribir({\n  documento: RegistroPage.documentoUnico(), edad: 17 });\n"
     "await registro.esperarResultado('Persona menor de edad');",
     size=11, highlight={0: C['blue600'], 1: C['blue600'], 2: C['blue600']})
# tres capas
capas = [('Pruebas', 'qué se prueba', C['blue100'], C['blue600']),
         ('RegistroPage', 'cómo se interactúa', C['plane'], C['ink']),
         ('La página', 'index.html', C['plane'], C['ink2'])]
for i, (a, b_, f, col) in enumerate(capas):
    bx = 6.9 + i * 2.0
    bb = tarjeta(s, bx, 3.75, 1.75, 0.95, fill=f)
    shape_text(bb, [[(a, {'bold': True, 'size': 13, 'color': col})], [(b_, {'size': 11, 'color': C['ink2']})]])
    if i < 2:
        flecha(s, bx + 1.75, 4.22, bx + 2.0, 4.22)
text(s, 0.6, 4.95, 12.1, 1.4, [
    [('Cuando la interfaz cambia, ', {'bold': True, 'color': C['ink']}),
     ('se corrige en un solo lugar y las pruebas siguen funcionando. Los campos que la prueba no declara '
      '(nombre, género, "está viva") toman valores por defecto razonables: cada prueba dice solo lo que le importa.', {})]],
     size=15, color=C['ink2'], line=1.2)


# ---------------------------------------------------------------- 8. prueba 09

s = d.slide('La prueba 09: una regla que la interfaz no deja alcanzar', 'Módulo 1 · La prueba que más enseña', notes=(
    'La prueba 09 prueba las dos capas. Por la interfaz es imposible provocar INVALID_AGE: el '
    'navegador bloquea 150 antes de enviar, y la prueba verifica que ni siquiera hubo petición. '
    'Si solo se prueba por la UI, se concluye que la regla del servidor no existe o sobra, y las '
    'dos conclusiones son falsas: la API se llama sin formulario, justo lo que hace el taller de carga. '
    'Por eso la entrega pide un escenario de este tipo.'))
navegador(s, 0.6, 1.55, 3.55, os.path.join(IMG, 'formulario-validacion.png'))
tarjeta(s, 4.6, 1.65, 8.15, 2.0, fill=C['blue100'])
text(s, 4.85, 1.82, 7.7, 0.35, 'Capa 1 · el navegador detiene el caso', size=15, bold=True, color=C['blue600'])
code(s, 4.85, 2.25, 7.7, 1.2,
     "await page.getByLabel('Edad').fill('150');\n...click();\n"
     "await expect(page.getByText('La edad debe estar entre 0 y 120.')).toBeVisible();\n"
     "expect(huboLlamada).toBe(false);   // nunca llegó al servidor",
     size=11, highlight={3: C['blue600']})
tarjeta(s, 4.6, 3.9, 8.15, 1.75, fill=C['orange100'])
text(s, 4.85, 4.07, 7.7, 0.35, 'Capa 2 · la API sí es alcanzable sin el formulario', size=15, bold=True, color=C['s2'])
code(s, 4.85, 4.5, 7.7, 0.95,
     "const r = await request.post('/register',\n  { data: { ..., age: 150 } });\n"
     "expect((await r.text()).trim()).toBe('INVALID_AGE');",
     size=11, highlight={2: C['s2']})
text(s, 4.6, 5.95, 8.15, 0.7, 'Una suite E2E mide lo que se alcanza a través de la interfaz. Eso no es lo mismo que lo que el sistema hace.',
     size=13, color=C['ink'], bold=True, line=1.15)


# ---------------------------------------------------------------- 9. accesibilidad

s = d.slide('Una suite de axe en verde no dice "el sitio es accesible"', 'Módulo 3 · Accesibilidad', notes=(
    'Contexto legal: el European Accessibility Act es exigible desde junio de 2025 y en Estados Unidos '
    'la ADA genera litigio constante. Pero lo importante del módulo es el limite de la herramienta: axe '
    'detecta de forma fiable cerca del 40 % de los problemas WCAG, los mecánicos. Lo demás exige juicio. '
    'El módulo audita también la página con errores visibles: los estados de error son el punto ciego clásico.'))
text(s, 0.6, 1.8, 4.2, 1.4, '≈ 40 %', size=72, bold=True, color=C['s1'], line=1.0)
text(s, 0.6, 3.05, 4.0, 1.2, 'de los problemas WCAG los detecta axe de forma fiable. El resto exige juicio humano.',
     size=15, color=C['ink2'], line=1.2)
text(s, 0.6, 4.3, 4.0, 1.6, [
    [('European Accessibility Act: ', {'bold': True, 'color': C['ink']}),
     ('exigible desde junio de 2025. En Estados Unidos, la ADA genera litigio constante sobre sitios web.', {})]],
     size=13, color=C['ink2'], line=1.2)
tarjeta(s, 5.2, 1.75, 3.65, 3.55, fill=C['blue100'])
text(s, 5.45, 1.95, 3.2, 0.4, 'Lo que mide una máquina', size=15, bold=True, color=C['blue600'])
tb = text(s, 5.45, 2.45, 3.2, 3.7, ['Contraste insuficiente', '<img> sin alt', 'Botones y enlaces sin nombre accesible',
                                    'Página sin idioma declarado', 'Un <select> sin nombre'],
          size=13, color=C['ink'], line=1.15)
para_space(tb, 10)
tarjeta(s, 9.1, 1.75, 3.65, 3.55, fill=C['orange100'])
text(s, 9.35, 1.95, 3.2, 0.4, 'Lo que exige juicio', size=15, bold=True, color=C['s2'])
tb = text(s, 9.35, 2.45, 3.2, 3.7, ['Si el texto alternativo es útil', 'Si el foco de teclado se ve',
                                    'Si un enlace dice a dónde lleva', 'Si un error explica qué corregir',
                                    'Si el resultado se anuncia'],
          size=13, color=C['ink'], line=1.15)
para_space(tb, 10)
text(s, 0.6, 5.75, 12.1, 0.8, [
    [('En verde significa otra cosa: ', {'bold': True, 'color': C['ink']}),
     ('"no tiene los errores que una máquina puede detectar sola". Por eso el módulo audita también la '
      'página con errores de validación visibles, el punto ciego clásico.', {})]],
     size=14, color=C['ink2'], line=1.2)


# ---------------------------------------------------------------- 10. 17 defectos

s = d.slide('17 defectos sembrados, en tres montones', 'Módulo 3B · Ver a la herramienta fallar', notes=(
    'defectuosa.html es la misma pantalla con 17 defectos puestos a mano. Lo que enseña el módulo es que '
    'no se reparten en dos montones sino en tres. A: axe los reporta con el filtro WCAG. B: solo si se quita '
    'el filtro. C: nunca, hay que buscarlos a mano. La frontera entre A y B la decide una línea de configuración. '
    'Las tres listas se obtuvieron ejecutando axe, no leyendo documentación: si una regla cambia de categoría, '
    'las pruebas del 3B fallan.'))
grupos = [('A', 7, 'axe los reporta con .withTags([...wcag...])', '6 reglas', C['s1'], C['blue100']),
          ('B', 3, 'axe los reporta solo si se quita ese filtro', '4 reglas', C['s2'], C['orange100']),
          ('C', 7, 'axe no los reporta nunca', 'revisión manual', C['ink2'], C['plane'])]
bx, bw_total, by = 0.6, 12.13, 1.85
unidad = bw_total / 17
for g, n, desc, cuenta, col, fnd in grupos:
    w_ = n * unidad
    b_ = box(s, bx, by, w_ - 0.06, 0.9, fill=col, radius=0.08)
    shape_text(b_, [[(f'Grupo {g}', {'bold': True, 'size': 15, 'color': '#ffffff'})],
                    [(f'{n} defectos', {'size': 13, 'color': '#ffffff'})]])
    bx += w_
reglas = {
    'A': ['button-name · color-contrast · html-has-lang', 'image-alt · link-name · select-name'],
    'B': ['heading-order', 'landmark-one-main', 'region', 'tabindex'],
    'C': 'placeholder como única etiqueta, alt="imagen", foco invisible, casilla "Estado", resultado sin role="status", '
         '"Haga clic aquí", un error que dice solo "Error."',
}
cx = 0.6
anchos = [7 * unidad - 0.06, 3 * unidad - 0.06, 7 * unidad - 0.06]
for (g, n, desc, cuenta, col, fnd), aw in zip(grupos, anchos):
    text(s, cx, 3.0, aw, 0.35, cuenta, size=13, bold=True, color=col)
    text(s, cx, 3.35, aw, 0.8, desc, size=13, color=C['ink'], line=1.15)
    text(s, cx, 4.15, aw, 1.7, reglas[g], size=12, color=C['ink2'], line=1.2,
         font=MONO if g != 'C' else FONT)
    cx += aw + 0.06
text(s, 0.6, 6.05, 12.1, 0.7, 'Axe encuentra 10 de 17. Ninguno de los 7 del grupo C: esos los localiza usted, a mano, y explica por qué ninguna herramienta puede detectarlos.',
     size=14, bold=True, line=1.15)


# ---------------------------------------------------------------- 11. el filtro

s = d.slide('Una línea de configuración esconde cuatro defectos reales', 'Módulo 3B · El filtro .withTags()', notes=(
    'Casi todos los tutoriales copian withTags con las etiquetas WCAG sin decir que ese filtro silencia la '
    'categoría best-practice de axe, donde viven el orden de encabezados, la ausencia de main y los tabindex '
    'positivos. La prueba 02 del módulo corre la misma página dos veces, con y sin filtro. '
    'Pregunta para el Wiki: cuál de las cuatro reglas escondidas le parece más grave y por qué.'))
text(s, 0.6, 1.7, 5.9, 0.4, 'Con filtro: 6 reglas', size=16, bold=True, color=C['s1'])
code(s, 0.6, 2.15, 5.9, 1.05, "new AxeBuilder({ page })\n  .withTags(['wcag2a', 'wcag2aa', 'wcag21a', 'wcag21aa'])\n  .analyze();",
     size=12, highlight={1: C['s1']})
text(s, 6.85, 1.7, 5.9, 0.4, 'Sin filtro: 10 reglas', size=16, bold=True, color=C['s2'])
code(s, 6.85, 2.15, 5.9, 1.05, "new AxeBuilder({ page })\n\n  .analyze();", size=12)
a_reglas = ['button-name', 'color-contrast', 'html-has-lang', 'image-alt', 'link-name', 'select-name']
b_reglas = ['heading-order', 'landmark-one-main', 'region', 'tabindex']
for i, r_ in enumerate(a_reglas):
    pill(s, 0.6 + (i % 2) * 2.95, 3.55 + (i // 2) * 0.5, r_, C['blue600'], C['blue100'], size=11, w=2.75)
for i, r_ in enumerate(a_reglas + b_reglas):
    oculto = r_ in b_reglas
    pill(s, 6.85 + (i % 2) * 2.95, 3.55 + (i // 2) * 0.5, r_,
         C['s2'] if oculto else C['blue600'], C['orange100'] if oculto else C['blue100'], size=11, w=2.75)
text(s, 6.85, 6.15, 5.9, 0.6, [[('En naranja: ', {'bold': True, 'color': C['s2']}),
                                ('la categoría best-practice, que el filtro silencia.', {})]],
     size=13, color=C['ink2'])
text(s, 0.6, 5.2, 5.9, 1.3, 'Los mismos defectos siguen en la página. Lo único que cambió es qué se le pidió a axe que mirara.',
     size=14, color=C['ink2'], line=1.2)


# ---------------------------------------------------------------- 12. placeholder

s = d.slide('Axe aprueba un campo sin etiqueta', 'Módulo 3B · Grupo C', notes=(
    'El defecto más instructivo del grupo C. El campo del nombre solo tiene placeholder, que desaparece al '
    'escribir. Axe lo da por bueno porque el placeholder cuenta como nombre accesible en el cálculo de accname. '
    'La única queja de axe sobre ese input es su tabindex positivo. La prueba 05 del módulo 3, escrita a mano, '
    'exige label for o aria-label y sí lo encuentra: esa es la razón concreta para escribir aserciones propias '
    'además de pasar la herramienta.'))
navegador(s, 0.6, 1.55, 4.55, os.path.join(IMG, 'defectuosa.png'), url='localhost:8080/defectuosa.html', max_h=5.3)
text(s, 5.65, 1.7, 7.1, 1.4, [
    [('El placeholder cuenta como nombre accesible ', {'bold': True, 'color': C['ink']}),
     ('en el cálculo de accname, así que la regla label da el campo por bueno. Pero desaparece en cuanto la persona empieza a escribir.', {})]],
     size=15, color=C['ink2'], line=1.2)
text(s, 5.65, 3.05, 7.1, 0.35, 'Los 7 defectos que ninguna herramienta ve', size=14, bold=True)
grupo_c = ['placeholder usado como única etiqueta', 'alt="imagen": existe, pero no describe nada',
           'outline: none, el foco de teclado se vuelve invisible', 'una casilla rotulada "Estado"',
           'el resultado sin role="status": nunca se anuncia', 'un enlace que dice "Haga clic aquí"',
           'un error que dice "Error." y nada más (WCAG 3.3.3)']
for i, t in enumerate(grupo_c):
    yy = 3.5 + i * 0.4
    dot(s, 5.8, yy + 0.14, 0.14, C['critical'] if i == 0 else C['axis'])
    text(s, 6.05, yy, 6.7, 0.35, t, size=13, color=C['ink'] if i == 0 else C['ink2'], bold=i == 0)
text(s, 5.65, 6.4, 7.1, 0.5, 'La prueba 05 del módulo 3, escrita a mano, sí lo encuentra.', size=13, bold=True, color=C['s1'])


# ---------------------------------------------------------------- 13. regresion visual

s = d.slide('Lo que ninguna aserción sobre el DOM ve', 'Módulo 4 · Regresión visual', notes=(
    'Las tres capturas de referencia reales del taller: el formulario vacío, un resultado de rechazo y el '
    'formulario en móvil. toHaveScreenshot compara píxel a píxel. El riesgo del patrón es actualizar las '
    'referencias sin mirar el diff: ahí la prueba deja de proteger. Las capturas dependen del sistema operativo '
    'y las fuentes, por eso el CI del taller no ejecuta este módulo.'))
text(s, 0.6, 1.62, 4.8, 0.3, 'formulario-vacio · 1280 px', size=11, color=C['muted'])
p1 = s.shapes.add_picture(os.path.join(SNAP, 'formulario-vacio-chromium-win32.png'), Inches(0.6), Inches(1.95), width=Inches(4.8))
p1.line.color.rgb = rgb(C['border'])
text(s, 0.6, 5.45, 4.8, 0.3, 'resultado-rechazo · solo el mensaje', size=11, color=C['muted'])
p2 = s.shapes.add_picture(os.path.join(SNAP, 'resultado-rechazo-chromium-win32.png'), Inches(0.6), Inches(5.75), width=Inches(4.8))
p2.line.color.rgb = rgb(C['border'])
text(s, 5.75, 1.62, 1.9, 0.3, 'móvil · 375 px', size=11, color=C['muted'])
p3 = s.shapes.add_picture(os.path.join(SNAP, 'formulario-movil-chromium-win32.png'), Inches(5.75), Inches(1.95), height=Inches(4.75))
p3.line.color.rgb = rgb(C['border'])
code(s, 8.0, 1.95, 4.75, 0.95, "npm run test:visual     # comparar\nnpm run visual:update   # actualizar",
     size=12)
text(s, 8.0, 3.15, 4.75, 1.4, [
    [('El riesgo: ', {'bold': True, 'color': C['critical']}),
     ('actualizar las referencias sin mirar el diff. Desde ese momento la prueba aprueba cualquier cosa.', {})]],
     size=14, color=C['ink2'], line=1.2)
text(s, 8.0, 4.6, 4.75, 1.6, [
    [('Dependen del sistema operativo: ', {'bold': True, 'color': C['ink']}),
     ('una referencia de Windows no coincide con la de un runner Linux. Por eso el CI del taller no ejecuta este módulo.', {})]],
     size=14, color=C['ink2'], line=1.2)


# ---------------------------------------------------------------- 14. usabilidad

s = d.slide('Tareas como objetivos, no como instrucciones', 'Módulo 5 · Usabilidad con usuarios', notes=(
    'El único módulo que no se automatiza. Cinco participantes bastan: con 5 personas se detecta cerca del '
    '85 % de los problemas (Nielsen); después cada uno repite hallazgos. Deben ser personas que no conozcan el '
    'proyecto. La diferencia entre instrucción y objetivo decide si la prueba sirve: la instrucción ya contiene '
    'la respuesta. Reglas del moderador: no ayudar, no explicar la interfaz, pedir que piensen en voz alta y '
    'registrar lo que hacen, no solo lo que dicen.'))
text(s, 0.6, 1.75, 3.3, 1.2, '5', size=80, bold=True, color=C['s2'], line=1.0)
text(s, 1.55, 1.95, 2.5, 1.0, 'participantes detectan cerca del 85 % de los problemas', size=14, color=C['ink2'], line=1.15)
text(s, 0.6, 3.2, 3.4, 0.4, 'Nielsen. Personas que no hayan trabajado en el proyecto.', size=11, color=C['muted'], line=1.15)
tarjeta(s, 4.4, 1.75, 4.05, 1.85, fill='#f6d8d8')
text(s, 4.65, 1.92, 3.6, 0.3, 'INSTRUCCIÓN · NO SIRVE', size=11, bold=True, color=C['critical'], spacing=1)
text(s, 4.65, 2.3, 3.6, 1.2, '"Escribe 30 en el campo Edad y pulsa Registrar"', size=15, line=1.2)
tarjeta(s, 8.7, 1.75, 4.05, 1.85, fill='#d7efd7')
text(s, 8.95, 1.92, 3.6, 0.3, 'OBJETIVO · SIRVE', size=11, bold=True, color=C['goodtext'], spacing=1)
text(s, 8.95, 2.3, 3.6, 1.2, '"Inscribe a tu tía Ana, de 30 años, para que pueda votar"', size=15, line=1.2)
text(s, 4.4, 3.72, 8.3, 0.6, 'La primera ya contiene la respuesta: solo comprueba que los dedos funcionan. La segunda revela si la persona encuentra el camino.',
     size=12, color=C['ink2'], line=1.15)
metricas = [('Tasa de éxito', '% que la completa sin ayuda'), ('Tiempo en tarea', 'segundos hasta completarla'),
            ('Errores', 'acciones equivocadas antes de acertar'), ('Abandono', '% que se rinde')]
for i, (m, desc) in enumerate(metricas):
    mx = 0.6 + i * 3.1
    tarjeta(s, mx, 4.65, 2.85, 1.1)
    text(s, mx + 0.22, 4.8, 2.5, 0.35, m, size=14, bold=True)
    text(s, mx + 0.22, 5.15, 2.5, 0.5, desc, size=12, color=C['ink2'], line=1.1)
text(s, 0.6, 6.05, 12.1, 0.8, [
    [('"No ayudes." ', {'bold': True, 'color': C['s2']}),
     ('Si alguien no encuentra el botón en 30 segundos, ese silencio incómodo es el hallazgo. Pida pensar en voz alta y anote lo que hacen, no solo lo que dicen.', {})]],
     size=13, color=C['ink2'], line=1.15)


# ---------------------------------------------------------------- 15. SUS

s = d.slide('SUS: 68 no es "68 %", es exactamente el promedio', 'Módulo 5 · System Usability Scale', notes=(
    'Diez preguntas al terminar la sesión, de 1 a 5. Las impares suman x - 1, las pares 5 - x, y la suma se '
    'multiplica por 2,5. El resultado va de 0 a 100 y no es un porcentaje. 68 es la media de la industria: la '
    'mitad de los sistemas medidos puntúan por debajo. Bandas de Sauro y Lewis.'))
text(s, 0.6, 1.75, 4.6, 0.4, 'Cómo se calcula', size=16, bold=True)
calc = [('1', 'Impares (1, 3, 5, 7, 9): ', 'x − 1'), ('2', 'Pares (2, 4, 6, 8, 10): ', '5 − x'),
        ('3', 'Sume los 10 valores y ', 'multiplique por 2,5')]
for i, (n, a, b_) in enumerate(calc):
    yy = 2.3 + i * 0.75
    dot(s, 0.8, yy + 0.2, 0.38, C['s1'])
    text(s, 0.61, yy + 0.06, 0.38, 0.3, n, size=13, bold=True, color='#ffffff', align=PP_ALIGN.CENTER)
    text(s, 1.2, yy + 0.04, 4.2, 0.6, [[(a, {}), (b_, {'bold': True, 'font': MONO, 'color': C['ink']})]],
         size=14, color=C['ink2'])
text(s, 0.6, 4.65, 4.5, 1.2, 'Resultado de 0 a 100. No es un porcentaje: es una posición frente a otros sistemas medidos.',
     size=14, color=C['ink2'], line=1.2)
# escala 0-100 con bandas
ex, ew, ey = 5.8, 6.9, 2.6
def xs(v):
    return ex + ew * v / 100
bandas = [(0, 51, 'Deficiente', '#f6d8d8', C['critical']), (51, 68, 'Aceptable', '#fdecc4', '#7a5600'),
          (68, 80.3, 'Bueno', C['blue100'], C['blue600']), (80.3, 100, 'Excelente', '#d7efd7', C['goodtext'])]
for a, b_, et, f, col in bandas:
    box(s, xs(a), ey, xs(b_) - xs(a) - 0.03, 1.0, fill=f, radius=0.0, shape=MSO_SHAPE.RECTANGLE)
    text(s, xs(a), ey + 0.32, xs(b_) - xs(a), 0.4, et, size=13, bold=True, color=col, align=PP_ALIGN.CENTER)
for v in [0, 51, 68, 80.3, 100]:
    etq = str(v).replace('.', ',')
    text(s, xs(v) - 0.5, ey + 1.1, 1.0, 0.3, etq, size=12, color=C['ink2'], align=PP_ALIGN.CENTER)
line(s, xs(68), ey - 0.35, xs(68), ey + 1.0, color=C['ink'], width=2)
text(s, xs(68) - 1.6, ey - 0.75, 3.2, 0.35, 'Promedio de la industria', size=13, bold=True, align=PP_ALIGN.CENTER)
text(s, ex, 4.25, ew, 1.6, [
    [('Ejemplo: ', {'bold': True, 'color': C['ink']}),
     ('un promedio de 72,5 queda en "bueno", apenas por encima de la media. Un 50 no es "la mitad de bien": '
      'es peor que la mayoría de los sistemas medidos.', {})]], size=14, color=C['ink2'], line=1.2)
text(s, ex, 5.75, ew, 0.4, 'Bandas de Sauro y Lewis. Más de 80,3 corresponde al percentil 90.', size=11, color=C['muted'])


# ---------------------------------------------------------------- 16. modulo 6: ciclo

s = d.slide('A la IA no se le cree: se la evalúa', 'Módulo 6 · Probar con IA · exploratorio', notes=(
    'Modulo exploratorio, no suma puntos. No se discute si la IA es buena: se mide, sobre esta aplicación, qué '
    'encuentra, qué se inventa y qué se le escapa. Los cuatro experimentos siguen el mismo ciclo, y el paso 4 es '
    'posible porque aquí tenemos algo que casi nunca hay en un proyecto real: una verdad conocida. Sabemos qué 17 '
    'defectos tiene defectuosa.html y qué reglas aplica la Registraduría, en qué orden. '
    'Herramientas: ChatGPT (responde lo que se le pega) y Claude Code (trabaja en una carpeta, ejecuta pruebas y '
    'maneja un navegador con Playwright MCP). Esa diferencia es una variable del experimento.'))
ciclo = [('1', 'Prompt', 'plantilla del taller'), ('2', 'Ejecutar', 'ChatGPT o Claude Code'),
         ('3', 'Salida', 'JSON o código'), ('4', 'Evaluar', 'script contra una verdad conocida'),
         ('5', 'Reporte', 'HTML que explica qué pasó')]
cw5, cg5 = 2.15, 0.345
for i, (n, t, desc) in enumerate(ciclo):
    cx5 = 0.6 + i * (cw5 + cg5)
    clave = i == 3
    b_ = tarjeta(s, cx5, 1.8, cw5, 1.45, fill=C['s7'] if clave else C['plane'])
    text(s, cx5 + 0.2, 1.95, 0.5, 0.4, n, size=20, bold=True, color='#ffffff' if clave else C['s7'])
    text(s, cx5 + 0.2, 2.38, cw5 - 0.35, 0.35, t, size=16, bold=True, color='#ffffff' if clave else C['ink'])
    text(s, cx5 + 0.2, 2.75, cw5 - 0.35, 0.45, desc, size=11, color='#ffffff' if clave else C['ink2'], line=1.1)
    if i < 4:
        flecha(s, cx5 + cw5 + 0.03, 2.52, cx5 + cw5 + cg5 - 0.03, 2.52)
exps = [('E1', 'Auditoría', '¿Encuentra la IA lo que axe no ve? Los 7 del grupo C.'),
        ('E4', 'Estabilidad', '¿Encuentra lo mismo en tres ejecuciones con el mismo prompt?'),
        ('E2', 'Pruebas generadas', '¿Detectan algo cuando se rompe la aplicación a propósito?'),
        ('E3', 'Plan exploratorio', '¿Qué reglas descubre un agente usando la app? ¿Caja negra o blanca?')]
for i, (e, t, desc) in enumerate(exps):
    ex6 = 0.6 + i * 3.1
    tarjeta(s, ex6, 3.65, 2.85, 1.85, fill='#ffffff', line_=C['border'])
    pill(s, ex6 + 0.2, 3.85, e, C['s7'], '#e4e0f6', size=11, w=0.6)
    text(s, ex6 + 0.95, 3.87, 1.8, 0.3, t, size=13, bold=True)
    text(s, ex6 + 0.2, 4.35, 2.45, 1.1, desc, size=12, color=C['ink2'], line=1.15)
text(s, 0.6, 5.8, 12.1, 0.9, [
    [('Sesión 1: ', {'bold': True, 'color': C['ink']}), ('E1 y E4, auditar defectuosa.html.   ', {}),
     ('Sesión 2: ', {'bold': True, 'color': C['ink']}), ('E2 y E3, generar pruebas y diseñar un plan.   ', {}),
     ('No suma puntos en la rúbrica.', {'bold': True, 'color': C['s7']})]], size=13, color=C['ink2'], line=1.2)


# ---------------------------------------------------------------- 17. la trampa

s = d.slide('Si la IA puede ver las respuestas, no las encuentra: las copia', 'Módulo 6 · La trampa que invalida el experimento',
            notes=(
    'El error más fácil de cometer. El taller tiene las respuestas en tres lugares: los comentarios de '
    'defectuosa.html, las pruebas de los módulos 1 y 2, y el README. Si la IA los ve, el resultado sale '
    'espectacular y no significa nada. En aprendizaje automático esto se llama evaluar con datos que el modelo '
    'ya vio. Por eso los scripts preparan copias limpias fuera del repositorio.'))
trampas = [('Los comentarios de defectuosa.html y el módulo 3B', 'E1 · E4', 'npm run ia:preparar',
            'copia sin comentarios, fuera del repositorio'),
           ('Las pruebas de los módulos 1 y 2', 'E2', 'npm run ia:preparar-generacion', 'una carpeta sin pruebas'),
           ('El README y las pruebas del taller', 'E3', 'caja negra: carpeta vacía', 'caja blanca: solo registraduria/')]
text(s, 0.6, 1.72, 4.5, 0.3, 'DÓNDE ESTÁN LAS RESPUESTAS', size=11, bold=True, color=C['muted'], spacing=1)
text(s, 5.35, 1.72, 1.5, 0.3, 'ARRUINA', size=11, bold=True, color=C['muted'], spacing=1)
text(s, 7.1, 1.72, 5.6, 0.3, 'CÓMO SE EVITA', size=11, bold=True, color=C['muted'], spacing=1)
for i, (donde, exp, cmd, desc) in enumerate(trampas):
    ty_ = 2.1 + i * 1.25
    tarjeta(s, 0.6, ty_, 12.15, 1.05)
    text(s, 0.85, ty_ + 0.15, 4.3, 0.8, donde, size=14, bold=True, line=1.15, anchor=MSO_ANCHOR.MIDDLE)
    pill(s, 5.35, ty_ + 0.36, exp, C['critical'], '#f6d8d8', size=11, w=1.2)
    text(s, 7.1, ty_ + 0.17, 5.5, 0.4, cmd, size=13, font=MONO if cmd.startswith('npm') else FONT, color=C['s7'], bold=True)
    text(s, 7.1, ty_ + 0.55, 5.5, 0.4, desc, size=12, color=C['ink2'])
text(s, 0.6, 6.0, 12.1, 0.8, 'En aprendizaje automático esto tiene nombre: evaluar un modelo con datos que ya vio. El resultado siempre sale espectacular y nunca significa nada.',
     size=14, color=C['ink2'], line=1.2)


# ---------------------------------------------------------------- 18. sabotajes

SUITE_TALLER = ['si', 'si', 'si', 'si', 'si', 'no']
SUITE_EJEMPLO = ['si', 'si', 'aviso', 'aviso', 'si', 'no']
s = d.slide('Que una prueba pase no dice si protege', 'Módulo 6 · E2, sabotajes', notes=(
    'Un proxy se pone delante de la Registraduría y le introduce seis defectos, uno a la vez. Cada suite se '
    'ejecuta siete veces: una limpia y una por sabotaje, y cada detección se confirma repitiendo. '
    'Detectado: alguna prueba falló. Sobrevivió: el sabotaje actuo y ninguna prueba lo notó. No ejercitado: '
    'ninguna prueba paso por ahí. La suite del taller detecta S4 porque la prueba 09 llama a la API a propósito, '
    'y no detecta S6 porque eso lo revisa el módulo 3, que no participa. El ejemplo está escrito a mano (no lo '
    'generó ninguna IA) e imita los defectos típicos de las pruebas generadas: pasa en verde con la aplicación '
    'intacta y solo detecta 3 de 6. Aclararlo en clase: no es un resultado de ChatGPT ni de Claude Code. '
    'Cifras de npm run ia:sabotaje -- --ejemplo.'))
sabs = ['S1 · Se inscriben menores de edad', 'S2 · Se aceptan documentos repetidos',
        'S3 · Se inscriben personas fallecidas', 'S4 · La API acepta edades imposibles',
        'S5 · Desaparece la validación del navegador', 'S6 · El resultado deja de anunciarse']
etq = {'si': 'Detectado', 'no': 'Sobrevivió', 'aviso': 'No ejercitado'}
text(s, 6.75, 1.68, 2.3, 0.35, 'Suite del taller', size=13, bold=True, align=PP_ALIGN.CENTER)
text(s, 9.1, 1.68, 2.3, 0.35, 'Ejemplo con defectos típicos', size=12, bold=True, align=PP_ALIGN.CENTER)
for i, nom in enumerate(sabs):
    ry = 2.1 + i * 0.58
    if i % 2 == 0:
        box(s, 0.6, ry - 0.06, 11.0, 0.52, fill=C['plane'], radius=0.1)
    text(s, 0.8, ry + 0.05, 5.8, 0.35, nom, size=14)
    estado(s, 6.75 + 0.37, ry + 0.04, etq[SUITE_TALLER[i]], SUITE_TALLER[i])
    estado(s, 9.1 + 0.37, ry + 0.04, etq[SUITE_EJEMPLO[i]], SUITE_EJEMPLO[i])
tot_t = SUITE_TALLER.count('si')
tot_e = SUITE_EJEMPLO.count('si')
text(s, 6.75, 5.62, 2.3, 0.5, f'{tot_t} de 6', size=24, bold=True, color=C['goodtext'], align=PP_ALIGN.CENTER)
text(s, 9.1, 5.62, 2.3, 0.5, f'{tot_e} de 6', size=24, bold=True, color=C['critical'], align=PP_ALIGN.CENTER)
text(s, 0.8, 5.65, 5.6, 1.2, [
    [('Sobrevivió ', {'bold': True, 'color': C['critical']}), ('es una aserción débil. ', {}),
     ('No ejercitado ', {'bold': True, 'color': '#7a5600'}), ('es un caso que no existe. ', {}),
     ('Las dos suites pasan en verde con la aplicación intacta.', {'bold': True, 'color': C['ink']})]],
     size=13, color=C['ink2'], line=1.2)


# ---------------------------------------------------------------- 19. playwright vs selenium

s = d.slide('Playwright o Selenium: la misma aplicación, dos pistas', 'Herramientas', notes=(
    'Playwright es la pista principal (28 pruebas) y Selenium con Java la alternativa (7). Las diferencias que '
    'importan en este taller: Playwright espera solo y levanta la aplicación antes de probar; con Selenium hay que '
    'esperar a mano con WebDriverWait y tener la app corriendo. Accesibilidad y regresión visual son nativas o casi '
    'en Playwright. Selenium sigue siendo el estándar histórico y soporta prácticamente todos los lenguajes.'))
filas = [('', 'Playwright', 'Selenium'),
         ('Espera automática', 'Sí', 'Manual (WebDriverWait)'),
         ('Levanta la aplicación', 'Sí (webServer)', 'No'),
         ('Accesibilidad', '@axe-core/playwright', 'Requiere integración'),
         ('Regresión visual', 'Nativa', 'Librería externa'),
         ('Lenguajes', 'JS/TS, Python, Java, .NET', 'Prácticamente todos'),
         ('Industria', 'En crecimiento', 'Estándar histórico'),
         ('En este taller', 'Pista principal · 28 pruebas', 'Pista alternativa · 7 pruebas')]
tab = s.shapes.add_table(len(filas), 3, Inches(0.6), Inches(1.75), Inches(12.13), Inches(0.52 * len(filas))).table
tab.columns[0].width = Inches(3.4)
tab.columns[1].width = Inches(4.37)
tab.columns[2].width = Inches(4.36)
tbl_pr = tab._tbl.tblPr
for attr in ('firstRow', 'bandRow'):
    tbl_pr.set(attr, '0')
for r_i, fila in enumerate(filas):
    for c_i, val in enumerate(fila):
        celda = tab.cell(r_i, c_i)
        celda.fill.solid()
        celda.fill.fore_color.rgb = rgb(C['plane'] if r_i % 2 == 0 and r_i > 0 else C['bg'])
        celda.margin_left = Inches(0.18)
        celda.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = celda.text_frame
        tf.paragraphs[0].text = ''
        run = tf.paragraphs[0].add_run()
        run.text = val
        run.font.name = FONT
        run.font.size = Pt(13 if r_i else 15)
        run.font.bold = r_i == 0 or c_i == 0
        run.font.color.rgb = rgb(C['s1'] if (r_i == 0 and c_i == 1) else C['ink'] if r_i == 0 or c_i == 0 else C['ink2'])


# ---------------------------------------------------------------- 20. entrega y rubrica

s = d.slide('Cómo se califica la entrega', 'Entrega', notes=(
    'La rúbrica suma 65 puntos: 13 criterios de 5 puntos. Son 10 filas, pero E2E, Accesibilidad y Usabilidad '
    'con usuarios valen por dos: son el núcleo del taller. La prueba definitiva antes de entregar: clonar el propio '
    'repositorio en otra carpeta y ejecutar npm test. El módulo 6 no suma puntos.'))
crit = [('Pruebas E2E', 10), ('Accesibilidad (axe / WCAG)', 10), ('Usabilidad con usuarios', 10),
        ('Estructura y ejecución', 5), ('Localizadores y esperas', 5), ('Page Object Model', 5),
        ('Regresión visual', 5), ('Gestión de defectos', 5), ('Integración continua', 5),
        ('Documentación y reflexión', 5)]
crit_inv = list(reversed(crit))
pl = Plot(s, 0.45, 1.6, 6.9, 5.2, XL_CHART_TYPE.BAR_CLUSTERED,
          cat([c for c, _ in crit_inv], [('Puntos', [v for _, v in crit_inv])]),
          yr=(0, 10), ymajor=5, inner=(0.42, 0.02, 0.52, 0.9), font_size=12, gap=45)
pl.style_series(0, C['s1'])
ser = pl.chart.plots[0].series[0]
for i, (_, v) in enumerate(crit_inv):
    if v == 10:
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(C['blue600'])
pl.chart.plots[0].has_data_labels = True
dl = pl.chart.plots[0].data_labels
dl.font.size = Pt(12)
dl.font.color.rgb = rgb(C['ink2'])
dl.position = XL_LABEL_POSITION.OUTSIDE_END
pl.chart.value_axis.visible = False
pl.chart.value_axis.has_major_gridlines = False
text(s, 7.8, 1.7, 5.0, 0.9, '65 puntos', size=40, bold=True, color=C['blue600'])
text(s, 7.8, 2.5, 5.0, 0.7, '13 criterios × 5. E2E, accesibilidad y usabilidad valen por dos.', size=13, color=C['ink2'], line=1.15)
bandas_r = [('59 – 65', 'Excelente dominio técnico y metodológico'), ('46 – 58', 'Buen trabajo, análisis parcial'),
            ('39 – 45', 'Cumple lo básico, sin profundidad'), ('< 39', 'No cumple los mínimos')]
for i, (rng, desc) in enumerate(bandas_r):
    yy = 3.35 + i * 0.47
    text(s, 7.8, yy, 1.2, 0.35, rng, size=13, bold=True)
    text(s, 9.05, yy, 3.7, 0.35, desc, size=12, color=C['ink2'])
tarjeta(s, 7.8, 5.35, 4.95, 1.3, fill=C['blue100'])
text(s, 8.05, 5.5, 4.5, 1.1, [
    [('Antes de entregar: ', {'bold': True, 'color': C['blue600']}),
     ('clone su propio repositorio en otra carpeta y ejecute npm test. Si no queda en verde sin pasos manuales, falta algo.', {})]],
     size=12, color=C['ink'], line=1.15)


# ---------------------------------------------------------------- 21. cierre

s = d.slide(notes=(
    'La pregunta que cierra el taller y que deben responder en el Wiki. Si la respuesta es "ninguna", lo más probable '
    'es que la sesión con usuarios se haya hecho con compañeros que ya conocian el formulario, o con tareas '
    'escritas como instrucciones. Las otras dos preguntas de la reflexión final están abajo.'))
text(s, 0.6, 1.3, 11, 0.3, 'LA PREGUNTA QUE CIERRA EL TALLER', size=12, bold=True, color=C['s2'], spacing=1.5)
text(s, 0.6, 1.75, 11.6, 2.4,
     '¿Qué problema encontraron los usuarios que ninguna de las 28 pruebas automatizadas podía detectar?',
     size=36, bold=True, line=1.1)
preguntas = ['¿Qué encontró axe que usted no habría notado mirando la pantalla?',
             '¿Qué le costó más: escribir las pruebas o mantenerlas estables?',
             'Si hizo el módulo 6: ¿para qué tarea usaría una IA al probar, para cuál no, y qué verificación manual no se saltaría nunca?']
for i, p_ in enumerate(preguntas):
    yy = 4.55 + i * 0.62
    dot(s, 0.75, yy + 0.16, 0.16, C['s1'])
    text(s, 1.05, yy, 11.5, 0.55, p_, size=16, color=C['ink2'], line=1.15)


d.save(SALIDA)
print(f'Presentacion generada: {SALIDA} ({d.n} diapositivas)')
